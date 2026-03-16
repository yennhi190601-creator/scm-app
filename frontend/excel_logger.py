import os
import openpyxl
from threading import Lock

EXCEL_PATH = "data/SCM_log.xlsx"

excel_lock = Lock()


def init_excel():

    os.makedirs("data", exist_ok=True)

    if not os.path.exists(EXCEL_PATH):

        wb = openpyxl.Workbook()
        ws = wb.active

        ws.append([
            "STT",
            "Ma YC",
            "Ngay YC",
            "Muc dich NC",
            "Ten vat tu",
            "YC KT",
            "Ref Part No",
            "SL",
            "DVT",
            "Du an",
            "Nhom VT",
            "Phan loai",
            "PIC NC",
            "Phong ban",
            "PIC MH",
            "Thoi han",
            "TBP duyet",
            "Nguoi nhap",
            "Tao luc"
        ])

        wb.save(EXCEL_PATH)


def append_excel(row):

    with excel_lock:

        wb = openpyxl.load_workbook(EXCEL_PATH)
        ws = wb.active

        ws.append([
            row.get("stt"),
            row.get("ma_yeu_cau"),
            row.get("ngay_yeu_cau"),
            row.get("muc_dich"),
            row.get("ten_vat_tu"),
            row.get("yc_ky_thuat"),
            row.get("ref_part_no"),
            row.get("so_luong"),
            row.get("dvt"),
            row.get("du_an"),
            row.get("nhom_vat_tu"),
            row.get("phan_loai"),
            row.get("pic_nghien_cuu"),
            row.get("phong_ban"),
            row.get("pic_mua_hang"),
            row.get("thoi_han"),
            row.get("tbp_duyet"),
            row.get("submitted_by"),
            row.get("created_at")
        ])

        wb.save(EXCEL_PATH)




# import os
# import openpyxl
# from datetime import datetime
# from threading import Lock

# excel_lock = Lock()


# def get_excel_path():

#     now = datetime.now()

#     year = now.year
#     month = now.month

#     quarter = (month - 1) // 3 + 1

#     filename = f"SCM_log_{year}_Q{quarter}.xlsx"

#     os.makedirs("data", exist_ok=True)

#     return os.path.join("data", filename)


# def init_excel():

#     path = get_excel_path()

#     if not os.path.exists(path):

#         wb = openpyxl.Workbook()
#         ws = wb.active

#         ws.append([
#             "STT",
#             "Ma YC",
#             "Ngay YC",
#             "Muc dich NC",
#             "Ten vat tu",
#             "YC KT",
#             "Ref Part No",
#             "SL",
#             "DVT",
#             "Du an",
#             "Nhom VT",
#             "Phan loai",
#             "PIC NC",
#             "Phong ban",
#             "PIC MH",
#             "Thoi han",
#             "TBP duyet",
#             "Nguoi nhap",
#             "Tao luc"
#         ])

#         wb.save(path)


# def append_excel(row):

#     path = get_excel_path()

#     with excel_lock:

#         if not os.path.exists(path):
#             init_excel()

#         wb = openpyxl.load_workbook(path)
#         ws = wb.active

#         ws.append([
#             row.get("stt"),
#             row.get("ma_yeu_cau"),
#             row.get("ngay_yeu_cau"),
#             row.get("muc_dich"),
#             row.get("ten_vat_tu"),
#             row.get("yc_ky_thuat"),
#             row.get("ref_part_no"),
#             row.get("so_luong"),
#             row.get("dvt"),
#             row.get("du_an"),
#             row.get("nhom_vat_tu"),
#             row.get("phan_loai"),
#             row.get("pic_nghien_cuu"),
#             row.get("phong_ban"),
#             row.get("pic_mua_hang"),
#             row.get("thoi_han"),
#             row.get("tbp_duyet"),
#             row.get("submitted_by"),
#             row.get("created_at")
#         ])

#         wb.save(path)
