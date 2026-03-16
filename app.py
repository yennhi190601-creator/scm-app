"""
SCM Research Request Management System
Version: 3.0.0 - Cloud Edition (PostgreSQL + Railway)
"""

from fastapi import FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from typing import Optional, List
import json
import os
import io
from datetime import datetime
from excel_logger import init_excel, append_excel   #Thêm vào để excel tự xuất file

# ── Detect DB mode ─────────────────────────────────────────────────────────────
DATABASE_URL = os.environ.get("DATABASE_URL", "")
USE_POSTGRES = bool(DATABASE_URL)

if USE_POSTGRES:
    import psycopg2
    import psycopg2.extras
    if DATABASE_URL.startswith("postgres://"):
        DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)
else:
    import sqlite3

# ── Load config ─────────────────────────────────────────────────────────────────
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, "config.json")
DB_PATH = os.path.join(BASE_DIR, "data", "scm_data.db")

def load_config():
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return json.load(f)

CONFIG = load_config()
if not USE_POSTGRES:
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)

# ── DB helpers ──────────────────────────────────────────────────────────────────
def get_conn():
    if USE_POSTGRES:
        return psycopg2.connect(DATABASE_URL)
    else:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        return conn

def pg_sql(sql):
    """Chuyển ? sang %s cho PostgreSQL"""
    return sql.replace("?", "%s") if USE_POSTGRES else sql

def row_to_dict(row, cursor):
    cols = [d[0] for d in cursor.description]
    return dict(zip(cols, row))

# ── App ────────────────────────────────────────────────────────────────────────
app = FastAPI(title="SCM Request System", docs_url=None, redoc_url=None)
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

# ── DB Init ────────────────────────────────────────────────────────────────────
def init_db():
    conn = get_conn()
    cur = conn.cursor()
    if USE_POSTGRES:
        cur.execute("""
            CREATE TABLE IF NOT EXISTS requests (
                stt              SERIAL PRIMARY KEY,
                ma_yeu_cau       TEXT, ngay_yeu_cau TEXT, muc_dich TEXT,
                ten_vat_tu       TEXT, yc_ky_thuat  TEXT, ref_part_no TEXT,
                so_luong         TEXT, dinh_muc     TEXT, chuc_nang TEXT,
                dvt              TEXT, du_an        TEXT, nhom_vat_tu TEXT,
                vi_tri_bom       TEXT, reference    TEXT, phan_loai TEXT,
                pic_nghien_cuu   TEXT, phong_ban    TEXT, pic_mua_hang TEXT,
                yeu_cau_tim_kiem TEXT, thoi_han     TEXT,
                tbp_duyet        TEXT DEFAULT 'Chua duyet',
                submitted_by     TEXT,
                created_at       TIMESTAMP DEFAULT NOW()
            )
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS edit_requests (
                id           SERIAL PRIMARY KEY,
                stt_ref      INTEGER, field_name TEXT, old_value TEXT,
                new_value    TEXT, reason TEXT, requested_by TEXT,
                status       TEXT DEFAULT 'pending',
                reviewed_by  TEXT, reviewed_at TEXT,
                created_at   TIMESTAMP DEFAULT NOW()
            )
        """)
    else:
        cur.execute("""
            CREATE TABLE IF NOT EXISTS requests (
                stt              INTEGER PRIMARY KEY AUTOINCREMENT,
                ma_yeu_cau       TEXT, ngay_yeu_cau TEXT, muc_dich TEXT,
                ten_vat_tu       TEXT, yc_ky_thuat  TEXT, ref_part_no TEXT,
                so_luong         TEXT, dinh_muc     TEXT, chuc_nang TEXT,
                dvt              TEXT, du_an        TEXT, nhom_vat_tu TEXT,
                vi_tri_bom       TEXT, reference    TEXT, phan_loai TEXT,
                pic_nghien_cuu   TEXT, phong_ban    TEXT, pic_mua_hang TEXT,
                yeu_cau_tim_kiem TEXT, thoi_han     TEXT,
                tbp_duyet        TEXT DEFAULT 'Chua duyet',
                submitted_by     TEXT,
                created_at       TEXT DEFAULT (datetime('now','localtime'))
            )
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS edit_requests (
                id           INTEGER PRIMARY KEY AUTOINCREMENT,
                stt_ref      INTEGER, field_name TEXT, old_value TEXT,
                new_value    TEXT, reason TEXT, requested_by TEXT,
                status       TEXT DEFAULT 'pending',
                reviewed_by  TEXT, reviewed_at TEXT,
                created_at   TEXT DEFAULT (datetime('now','localtime'))
            )
        """)
    conn.commit()
    conn.close()

init_db()
init_excel()
# ── Auth ────────────────────────────────────────────────────────────────────────
def get_role(email: str) -> str:
    cfg = load_config()
    admins  = [e.strip().lower() for e in cfg.get("admin_emails", [])]
    allowed = [e.strip().lower() for e in cfg.get("allowed_emails", [])]
    em = email.strip().lower()
    if em in admins:  return "admin"
    # if not allowed or em in allowed: return "user"
    domain = em.split("@")[-1]
    if not allowed or em in allowed or domain in allowed: return "user"
    return "blocked"

def require_auth(request: Request):
    email = request.headers.get("X-User-Email", "").strip().lower()
    if not email: raise HTTPException(401, "Chua dang nhap")
    role = get_role(email)
    if role == "blocked": raise HTTPException(403, "Email khong duoc cap quyen")
    return email, role

def require_admin(request: Request):
    email, role = require_auth(request)
    if role != "admin": raise HTTPException(403, "Chi admin moi co quyen nay")
    return email

# ── Models ──────────────────────────────────────────────────────────────────────
class RequestItem(BaseModel):
    ma_yeu_cau: Optional[str] = ""; ngay_yeu_cau: Optional[str] = ""
    muc_dich: Optional[str] = "";   ten_vat_tu: Optional[str] = ""
    yc_ky_thuat: Optional[str] = "";ref_part_no: Optional[str] = ""
    so_luong: Optional[str] = "";   dinh_muc: Optional[str] = ""
    chuc_nang: Optional[str] = "";  dvt: Optional[str] = ""
    du_an: Optional[str] = "";      nhom_vat_tu: Optional[str] = ""
    vi_tri_bom: Optional[str] = ""; reference: Optional[str] = ""
    phan_loai: Optional[str] = "";  pic_nghien_cuu: Optional[str] = ""
    phong_ban: Optional[str] = "";  pic_mua_hang: Optional[str] = ""
    yeu_cau_tim_kiem: Optional[str] = ""; thoi_han: Optional[str] = ""
    tbp_duyet: Optional[str] = "Chua duyet"

# ── Routes ──────────────────────────────────────────────────────────────────────
@app.post("/api/auth/login")
async def login(request: Request):
    body = await request.json()
    email = body.get("email","").strip().lower()
    if not email: raise HTTPException(400, "Email khong hop le")
    role = get_role(email)
    if role == "blocked": raise HTTPException(403, "Email khong duoc cap quyen")
    return {"email": email, "role": role}

@app.get("/api/config")
def get_config_route(request: Request):
    require_auth(request)
    cfg = load_config()
    return {k: v for k, v in cfg.items() if k not in ("admin_emails","allowed_emails")}

@app.post("/api/requests")
async def create_requests(request: Request):
    email, _ = require_auth(request)
    body  = await request.json()
    items = body.get("items", [])
    conn  = get_conn()
    cur   = conn.cursor()
    inserted = []
    for item in items:
        params = (
            item.get("ma_yeu_cau",""),   item.get("ngay_yeu_cau",""),
            item.get("muc_dich",""),      item.get("ten_vat_tu",""),
            item.get("yc_ky_thuat",""),   item.get("ref_part_no",""),
            item.get("so_luong",""),      item.get("dinh_muc",""),
            item.get("chuc_nang",""),     item.get("dvt",""),
            item.get("du_an",""),         item.get("nhom_vat_tu",""),
            item.get("vi_tri_bom",""),    item.get("reference",""),
            item.get("phan_loai",""),     item.get("pic_nghien_cuu",""),
            item.get("phong_ban",""),     item.get("pic_mua_hang",""),
            item.get("yeu_cau_tim_kiem",""), item.get("thoi_han",""),
            item.get("tbp_duyet","Chua duyet"), email
        )
        if USE_POSTGRES:
            cur.execute("""
                INSERT INTO requests (ma_yeu_cau,ngay_yeu_cau,muc_dich,ten_vat_tu,yc_ky_thuat,
                ref_part_no,so_luong,dinh_muc,chuc_nang,dvt,du_an,nhom_vat_tu,vi_tri_bom,
                reference,phan_loai,pic_nghien_cuu,phong_ban,pic_mua_hang,yeu_cau_tim_kiem,
                thoi_han,tbp_duyet,submitted_by) VALUES %s RETURNING stt
            """ % ("(" + ",".join(["%s"]*22) + ")"), params)
            # inserted.append(cur.fetchone()[0])      #Code cũ lấy thẳng database
           # Code mới để append excel
            new_id = cur.fetchone()[0]
            inserted.append(new_id)
            item["stt"] = new_id
            item["submitted_by"] = email
            item["created_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            append_excel(item)
        else:
            cur.execute("""
                INSERT INTO requests (ma_yeu_cau,ngay_yeu_cau,muc_dich,ten_vat_tu,yc_ky_thuat,
                ref_part_no,so_luong,dinh_muc,chuc_nang,dvt,du_an,nhom_vat_tu,vi_tri_bom,
                reference,phan_loai,pic_nghien_cuu,phong_ban,pic_mua_hang,yeu_cau_tim_kiem,
                thoi_han,tbp_duyet,submitted_by) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """, params)
            # inserted.append(cur.lastrowid)    #Code cũ lấy thẳng database
            # Code mới để append excel
            new_id = cur.lastrowid
            inserted.append(new_id)
            item["stt"] = new_id
            item["submitted_by"] = email
            item["created_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            append_excel(item)
    
    conn.commit(); conn.close()
    return {"success": True, "inserted": len(inserted), "ids": inserted}

@app.get("/api/requests")
def get_requests(request: Request, page: int=1, page_size: int=50, search: str=""):
    require_auth(request)
    conn = get_conn(); cur = conn.cursor()
    offset = (page-1)*page_size
    if search:
        q = f"%{search}%"
        sql = pg_sql("SELECT * FROM requests WHERE ma_yeu_cau LIKE ? OR ten_vat_tu LIKE ? OR pic_nghien_cuu LIKE ? OR submitted_by LIKE ? ORDER BY stt DESC LIMIT ? OFFSET ?")
        cur.execute(sql, (q,q,q,q,page_size,offset))
    else:
        cur.execute(pg_sql("SELECT * FROM requests ORDER BY stt DESC LIMIT ? OFFSET ?"), (page_size,offset))
    rows = [row_to_dict(r,cur) for r in cur.fetchall()]
    cur.execute("SELECT COUNT(*) FROM requests")
    total = cur.fetchone()[0]
    conn.close()
    return {"data": rows, "total": total, "page": page, "page_size": page_size}

@app.patch("/api/requests/{stt}/duyet")
async def toggle_duyet(stt: int, request: Request):
    require_admin(request)
    body = await request.json()
    status = body.get("tbp_duyet","Chua duyet")
    conn = get_conn(); cur = conn.cursor()
    cur.execute(pg_sql("UPDATE requests SET tbp_duyet=? WHERE stt=?"), (status,stt))
    conn.commit(); conn.close()
    return {"success": True}

@app.post("/api/edit-requests")
async def submit_edit_request(request: Request):
    email, _ = require_auth(request)
    body = await request.json()
    stt_ref    = body.get("stt_ref")
    field_name = body.get("field_name","")
    old_value  = body.get("old_value","")
    new_value  = body.get("new_value","")
    reason     = body.get("reason","")
    conn = get_conn(); cur = conn.cursor()
    cur.execute(pg_sql("SELECT stt FROM requests WHERE stt=?"), (stt_ref,))
    if not cur.fetchone(): conn.close(); raise HTTPException(404,"Khong tim thay ban ghi")
    cur.execute(pg_sql("SELECT id FROM edit_requests WHERE stt_ref=? AND field_name=? AND status='pending'"), (stt_ref,field_name))
    if cur.fetchone(): conn.close(); raise HTTPException(400,"Da co yeu cau dang cho duyet")
    cur.execute(pg_sql("INSERT INTO edit_requests (stt_ref,field_name,old_value,new_value,reason,requested_by) VALUES (?,?,?,?,?,?)"),
                (stt_ref,field_name,old_value,new_value,reason,email))
    conn.commit(); conn.close()
    return {"success": True}

@app.get("/api/edit-requests")
def get_edit_requests(request: Request, status: str="pending"):
    require_admin(request)
    conn = get_conn(); cur = conn.cursor()
    cur.execute(pg_sql("SELECT * FROM edit_requests WHERE status=? ORDER BY created_at DESC"), (status,))
    rows = [row_to_dict(r,cur) for r in cur.fetchall()]
    conn.close()
    return {"data": rows}

@app.get("/api/edit-requests/my")
def get_my_edit_requests(request: Request):
    email, _ = require_auth(request)
    conn = get_conn(); cur = conn.cursor()
    cur.execute(pg_sql("SELECT * FROM edit_requests WHERE requested_by=? ORDER BY created_at DESC LIMIT 50"), (email,))
    rows = [row_to_dict(r,cur) for r in cur.fetchall()]
    conn.close()
    return {"data": rows}

@app.patch("/api/edit-requests/{edit_id}/review")
async def review_edit_request(edit_id: int, request: Request):
    admin_email = require_admin(request)
    body   = await request.json()
    action = body.get("action","")
    conn   = get_conn(); cur = conn.cursor()
    cur.execute(pg_sql("SELECT * FROM edit_requests WHERE id=?"), (edit_id,))
    row = cur.fetchone()
    if not row: conn.close(); raise HTTPException(404,"Khong tim thay yeu cau")
    edit = row_to_dict(row,cur)
    if edit["status"] != "pending": conn.close(); raise HTTPException(400,"Yeu cau nay da xu ly roi")
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    allowed_fields = ["ma_yeu_cau","ngay_yeu_cau","muc_dich","ten_vat_tu","yc_ky_thuat",
        "ref_part_no","so_luong","dinh_muc","chuc_nang","dvt","du_an","nhom_vat_tu",
        "vi_tri_bom","reference","phan_loai","pic_nghien_cuu","phong_ban","pic_mua_hang",
        "yeu_cau_tim_kiem","thoi_han"]
    if action == "approve":
        if edit["field_name"] not in allowed_fields:
            conn.close(); raise HTTPException(400,"Truong khong hop le")
        cur.execute(pg_sql(f"UPDATE requests SET {edit['field_name']}=? WHERE stt=?"),
                    (edit["new_value"],edit["stt_ref"]))
        new_status = "approved"
    else:
        new_status = "rejected"
    cur.execute(pg_sql("UPDATE edit_requests SET status=?,reviewed_by=?,reviewed_at=? WHERE id=?"),
                (new_status,admin_email,now,edit_id))
    conn.commit(); conn.close()
    return {"success": True, "action": new_status}

@app.get("/api/export")
# def export_excel(request: Request):
def export_excel(request: Request, email: str = ""):
    if email:
        role = get_role(email)
        if role != "admin":
            raise HTTPException(403, "Khong phai admin")
    else:
        require_admin(request)
    # require_admin(request)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(500,"openpyxl not installed")
    conn = get_conn(); cur = conn.cursor()
    cur.execute("SELECT * FROM requests ORDER BY stt ASC")
    rows = cur.fetchall()
    cols = [d[0] for d in cur.description]
    conn.close()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "SCM Requests"
    header_map = {
        "stt":"STT","ma_yeu_cau":"Ma yeu cau","ngay_yeu_cau":"Ngay yeu cau",
        "muc_dich":"Muc dich NC","ten_vat_tu":"Ten vat tu","yc_ky_thuat":"YC KT ban dau",
        "ref_part_no":"Ref Part No","so_luong":"So luong","dinh_muc":"Dinh muc",
        "chuc_nang":"Chuc nang SD","dvt":"DVT","du_an":"Du an SD",
        "nhom_vat_tu":"Nhom vat tu","vi_tri_bom":"Vi tri BOM","reference":"Reference",
        "phan_loai":"Phan loai","pic_nghien_cuu":"PIC NC","phong_ban":"Phong ban NC",
        "pic_mua_hang":"PIC mua hang","yeu_cau_tim_kiem":"YC tim kiem",
        "thoi_han":"Thoi han TK","tbp_duyet":"TBP duyet",
        "submitted_by":"Nguoi nhap","created_at":"Thoi gian tao"
    }
    thin = Side(style="thin",color="CCCCCC")
    border = Border(left=thin,right=thin,top=thin,bottom=thin)
    for ci,col in enumerate(cols,1):
        cell = ws.cell(row=1,column=ci,value=header_map.get(col,col))
        cell.fill = PatternFill("solid",fgColor="1a3a5c")
        cell.font = Font(bold=True,color="FFFFFF",size=10)
        cell.alignment = Alignment(horizontal="center",vertical="center",wrap_text=True)
        cell.border = border
    for ri,row in enumerate(rows,2):
        for ci,val in enumerate(row,1):
            cell = ws.cell(row=ri,column=ci,value=val)
            cell.alignment = Alignment(vertical="center",wrap_text=True)
            cell.border = border
            if ri%2==0: cell.fill = PatternFill("solid",fgColor="F0F4F8")
    for ci in range(1,len(cols)+1):
        ws.column_dimensions[ws.cell(row=1,column=ci).column_letter].width = 18
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    filename = f"SCM_Requests_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"})

# ── Serve frontend ──────────────────────────────────────────────────────────────
frontend_dir = os.path.join(BASE_DIR, "frontend")
if os.path.exists(frontend_dir):
    app.mount("/static", StaticFiles(directory=frontend_dir), name="static")

@app.get("/")
def serve_index():
    return FileResponse(os.path.join(BASE_DIR, "frontend", "index.html"))

if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8000))
    uvicorn.run("app:app", host="0.0.0.0", port=port, reload=False)
